# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x0B, 0x2B, 0x4D)
BLUE = RGBColor(0x1F, 0x6F, 0xC0)
LIGHT = RGBColor(0xF2, 0xF6, 0xFA)
GRAY = RGBColor(0x59, 0x5F, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xE8, 0x7A, 0x1E)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(blank)


def set_bg(slide, color=WHITE):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color, line=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line:
        shp.line.color.rgb = color
        shp.line.width = Pt(0.5)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, text, size=18, color=NAVY, bold=False,
             align=PP_ALIGN.LEFT, font="Microsoft YaHei", anchor=None, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = font
    return tb


def add_bullets(slide, x, y, w, h, items, size=15, color=GRAY, bold_first=False, line_space=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (lvl, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        p.line_spacing = line_space
        p.space_after = Pt(6)
        r = p.add_run()
        prefix = "• " if lvl == 0 else "- "
        r.text = prefix + text
        r.font.size = Pt(size - lvl * 1.5)
        r.font.color.rgb = NAVY if lvl == 0 else color
        r.font.bold = (lvl == 0)
        r.font.name = "Microsoft YaHei"
    return tb


def footer(slide, idx, total=10, title=""):
    add_rect(slide, 0, SH - Inches(0.32), SW, Inches(0.32), NAVY)
    add_text(slide, Inches(0.4), SH - Inches(0.34), Inches(6), Inches(0.3),
              "中长期交易基本说明与业务逻辑", size=10, color=WHITE)
    add_text(slide, SW - Inches(1.2), SH - Inches(0.34), Inches(0.8), Inches(0.3),
              f"{idx:02d} / {total}", size=10, color=WHITE, align=PP_ALIGN.RIGHT)


def title_bar(slide, text, subtitle=None):
    add_rect(slide, 0, 0, SW, Inches(1.15), NAVY)
    add_rect(slide, 0, Inches(1.15), SW, Inches(0.06), ORANGE)
    add_text(slide, Inches(0.6), Inches(0.18), Inches(11), Inches(0.6), text,
              size=28, color=WHITE, bold=True)
    if subtitle:
        add_text(slide, Inches(0.6), Inches(0.72), Inches(11), Inches(0.4), subtitle,
                  size=13, color=RGBColor(0xCF, 0xDD, 0xEC))


TOTAL = 10

# ---------- Slide 1: Cover ----------
s = add_slide()
set_bg(s, NAVY)
add_rect(s, 0, Inches(6.7), SW, Inches(0.05), ORANGE)
add_text(s, Inches(1), Inches(2.5), Inches(11.3), Inches(1.2),
          "中长期交易基本说明与业务逻辑", size=40, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(3.6), Inches(11.3), Inches(0.6),
          "广西电力市场 · 中长期交易规则与策略汇报", size=18, color=RGBColor(0xCF, 0xDD, 0xEC), align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(6.85), Inches(11.3), Inches(0.4),
          "V1.0.1", size=12, color=RGBColor(0x9A, 0xB3, 0xCE), align=PP_ALIGN.CENTER)

# ---------- Slide 2: 目录 ----------
s = add_slide()
set_bg(s)
title_bar(s, "目  录", "CONTENTS")
toc = [
    "01  关键名词与基本概念",
    "02  交易目标：补仓与套利",
    "03  交易品种总览与核心公式",
    "04  多日用电合同电量转让交易（上午）— 规则",
    "05  多日用电合同电量转让交易 — 交易策略",
    "06  多日直接交易（下午）— 规则",
    "07  多日直接交易 — 交易策略",
    "08  月度策略与全局套利构想（待讨论）",
    "09  补充内容与待完善方向",
]
y = Inches(1.6)
for i, item in enumerate(toc):
    add_rect(s, Inches(0.6), y, Inches(0.06), Inches(0.5), ORANGE if i % 2 == 0 else BLUE)
    add_text(s, Inches(0.85), y - Inches(0.02), Inches(11), Inches(0.5), item, size=17, color=NAVY)
    y += Inches(0.58)
footer(s, 2, TOTAL)

# ---------- Slide 3: 关键名词 ----------
s = add_slide()
set_bg(s)
title_bar(s, "关键名词解释", "01 基本概念")
terms = [
    ("交易日", "申报电量、电价的日期"),
    ("运行日（D日 / 交割日）", "实际运行的日期；如 3.1 日申报 3.4 日的中长期交易，3.1 为交易日，3.4 为运行日"),
    ("D-1 / D+1", "对应交割日的前一天 / 后一天，以此类推"),
    ("能量块交易", "本文交易均为 1 小时能量块交易；如“3.1日中长期交易”指当日 0:00~23:00 共 24 个时刻点的 24 次独立交易，各自独立预测、独立报量报价"),
]
y = Inches(1.55)
for t, d in terms:
    add_rect(s, Inches(0.6), y, Inches(12.1), Inches(1.0), LIGHT)
    add_rect(s, Inches(0.6), y, Inches(0.08), Inches(1.0), BLUE)
    add_text(s, Inches(0.95), y + Inches(0.08), Inches(2.6), Inches(0.8), t, size=16, color=NAVY, bold=True)
    add_text(s, Inches(3.7), y + Inches(0.08), Inches(8.7), Inches(0.85), d, size=13.5, color=GRAY)
    y += Inches(1.18)
footer(s, 3, TOTAL)

# ---------- Slide 4: 交易目标 ----------
s = add_slide()
set_bg(s)
title_bar(s, "交易目标：补仓与套利", "02 Trading Objectives")
cards = [
    ("补仓", BLUE,
     ["广西本年度中长期持仓偏低，目标补全 50%（未来可能调整）的中长期持仓（以月为单位）",
      "必须每日购买一定中长期电量",
      "该交易通常在早晨的多日合同转让市场中完成"]),
    ("套利", ORANGE,
     ["实时电价在特定时段存在波动（可高可低）",
      "实时电价 - 中长期电价 的差值在不同时段可正可负，形成买卖需求",
      "不同售电公司对实时价格判断不同，部分公司博弈高/低价，形成交易与套利空间"]),
]
x = Inches(0.6)
for name, color, items in cards:
    add_rect(s, x, Inches(1.5), Inches(5.95), Inches(0.7), color)
    add_text(s, x, Inches(1.55), Inches(5.95), Inches(0.6), name, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_rect(s, x, Inches(2.25), Inches(5.95), Inches(4.3), LIGHT)
    bl = [(0, it) for it in items]
    add_bullets(s, x + Inches(0.35), Inches(2.5), Inches(5.3), Inches(3.8), bl, size=15)
    x += Inches(6.2)
footer(s, 4, TOTAL)

# ---------- Slide 5: 交易品种总览 + 公式 ----------
s = add_slide()
set_bg(s)
title_bar(s, "交易品种总览与核心公式", "03 Overview & Key Formulas")
add_text(s, Inches(0.6), Inches(1.4), Inches(6), Inches(0.4), "核心持仓公式", size=16, bold=True, color=NAVY)
formulas = [
    "月度中长期持仓量 = 年度长协(分解到月) + 月度竞价 + 月内多日直接交易 + 合同电量转让 + 绿色电力交易",
    "总电量(批发侧) = 月度中长期持仓量 + 日前偏差电量 + 实时偏差电量",
    "总电量(代理用户侧) = 代理所有零售用户月度实际用网电量之和",
    "正常情况下：批发侧口径 = 代理用户侧口径（特殊计量/非市场化情形除外，暂不考虑）",
]
add_bullets(s, Inches(0.6), Inches(1.8), Inches(12.1), Inches(2.4), [(0, f) for f in formulas], size=13.5)

add_rect(s, Inches(0.6), Inches(4.35), Inches(12.1), Inches(0.05), ORANGE)
add_text(s, Inches(0.6), Inches(4.5), Inches(6), Inches(0.4), "两大交易品种", size=16, bold=True, color=NAVY)

x = Inches(0.6)
for name, time, color in [("多日用电合同电量转让交易", "上午 09:00–11:30", BLUE),
                           ("多日直接交易", "下午 16:00–17:30", ORANGE)]:
    add_rect(s, x, Inches(5.0), Inches(5.95), Inches(1.6), LIGHT)
    add_rect(s, x, Inches(5.0), Inches(5.95), Inches(0.08), color)
    add_text(s, x + Inches(0.3), Inches(5.2), Inches(5.4), Inches(0.5), name, size=17, bold=True, color=NAVY)
    add_text(s, x + Inches(0.3), Inches(5.75), Inches(5.4), Inches(0.6), f"申报时段：{time}", size=14, color=GRAY)
    x += Inches(6.2)
footer(s, 5, TOTAL)

# ---------- Slide 6: 合同转让 — 规则 ----------
s = add_slide()
set_bg(s)
title_bar(s, "多日用电合同电量转让交易（上午） — 规则", "04 Rules")
left_items = [
    (0, "政策依据"),
    (1, "《广西电力中长期交易规则》《2026年广西电力市场化交易工作有关事项的通知》"),
    (1, "标的物为当月内未履约的年度直接交易合同电量"),
    (0, "交易目标"),
    (1, "2–3天短周期，匹配短期负荷波动、降低偏差风险"),
    (1, "滚动衔接，确保全月有可交易标的"),
    (0, "交易时序"),
    (1, "09:00–09:20 集中竞价（自由报价，可改可撤）"),
    (1, "09:20–09:30 静默固化、集中统一边际出清"),
    (1, "09:30 公布出清结果"),
    (1, "09:30–11:30 滚动撮合阶段"),
]
add_bullets(s, Inches(0.6), Inches(1.45), Inches(6.4), Inches(5.6), left_items, size=13.5)

add_rect(s, Inches(7.25), Inches(1.45), Inches(5.5), Inches(5.5), LIGHT)
add_text(s, Inches(7.55), Inches(1.65), Inches(5), Inches(0.4), "特殊规则约束", size=15, bold=True, color=NAVY)
rule_items = [
    (0, "触发条件：月度签约持仓量(考核部分) < 实际用电量 × 50%"),
    (0, "考核电量 = 实际用电量×50% − 持仓量(考核部分)"),
    (0, "考核价格 = 336.56元/MWh × 1.05 − 当月发电侧实时市场加权均价"),
    (1, "价格为正：按该价×考核电量缴纳考核费用"),
    (1, "价格为负：按绝对值×考核电量执行考核（反向约束，防止投机性低持仓）"),
    (0, "持仓量(考核部分) = 年度长协(本月) + 合同电量转让电量(本月)"),
    (0, "电量上下限见 PM 网站「滚动撮合-时刻剩余可买入电量」"),
    (0, "交易方向：买方/卖方均可，对手方为省内售电公司"),
]
add_bullets(s, Inches(7.55), Inches(2.1), Inches(5.0), Inches(4.7), rule_items, size=12.5)
footer(s, 6, TOTAL)

# ---------- Slide 7: 合同转让 — 策略 ----------
s = add_slide()
set_bg(s)
title_bar(s, "多日用电合同电量转让交易 — 交易策略", "05 Strategy")
add_text(s, Inches(0.6), Inches(1.35), Inches(6), Inches(0.4), "核心决策逻辑", size=15, bold=True, color=NAVY)
add_rect(s, Inches(0.6), Inches(1.8), Inches(12.1), Inches(0.9), LIGHT)
add_text(s, Inches(0.85), Inches(1.92), Inches(11.6), Inches(0.7),
          "盈亏线 = 合同电量转让价格 − （现货价格_预测 + 考核价格）\n"
          "盈亏线 < 0 → 买入（主要方式）；盈亏线 > 0 → 卖出（持仓不足50%前提下）",
          size=14, color=NAVY, bold=True)

items = [
    (0, "集中竞价阶段：人工权衡出清概率与最大收益，一次性表格导入+部分手动调整；绝大部分电量在此阶段成交"),
    (0, "滚动撮合阶段：实时盯盘报价，但因市场原因成交量极少"),
    (0, "经验规则"),
    (1, "以买入补仓为主；晚高峰（17:00–22:00，受新能源影响）即便转让价高也不轻易卖出，因新能源限电易致价格跳变"),
    (1, "若部分时刻卖出，倾向在其他时刻“换仓”买入近似等量电量，保持总持仓目标不变"),
    (1, "12:00–16:00 风电大发时段，现货预测价可能为0；部分交易员仍以 30–50 元/MWh 价格“赌”限电导致的价格跳变"),
]
add_bullets(s, Inches(0.6), Inches(2.95), Inches(12.1), Inches(4.0), items, size=14)
footer(s, 7, TOTAL)

# ---------- Slide 8: 多日直接交易 — 规则 ----------
s = add_slide()
set_bg(s)
title_bar(s, "多日直接交易（下午） — 规则", "06 Rules")
left_items = [
    (0, "政策依据"),
    (1, "标的为当月内未来2–3天可直接交割的增量市场化电量（含发电侧新增上网、用户侧新增购电）"),
    (1, "面向发电交易单元、售电公司、批发直购用户、虚拟电厂等"),
    (0, "交易目标"),
    (1, "衔接月度长协与日前现货，新增锁定未来数日基础电量、补齐缺口、降低偏差考核风险"),
    (0, "交易时序"),
    (1, "16:00–16:20 集中竞价（自由报价，可改可撤）"),
    (1, "16:20–16:30 静默固化、集中统一边际出清"),
    (1, "16:30–17:30 滚动撮合阶段"),
    (1, "数据爬取：PM网站「集中竞价-滚动撮合-多日直接交易」栏目"),
]
add_bullets(s, Inches(0.6), Inches(1.45), Inches(6.4), Inches(5.6), left_items, size=13.5)

add_rect(s, Inches(7.25), Inches(1.45), Inches(5.5), Inches(3.3), LIGHT)
add_text(s, Inches(7.55), Inches(1.65), Inches(5), Inches(0.4), "特殊规则补充", size=15, bold=True, color=NAVY)
add_bullets(s, Inches(7.55), Inches(2.15), Inches(5.0), Inches(2.5),
            [(0, "出售量上限 = 月度竞价 + 绿电交易 之和"),
             (1, "例：年长协40% + 月竞10% + 绿电0%，合同转让买入15%后总持仓达55%/65%，但多日直接交易卖出上限仍为10%（不受年度长协/合同转让影响）")],
            size=12.5)

add_rect(s, Inches(7.25), Inches(5.0), Inches(5.5), Inches(1.95), LIGHT)
add_text(s, Inches(7.55), Inches(5.2), Inches(5), Inches(0.4), "市场特点", size=15, bold=True, color=NAVY)
add_bullets(s, Inches(7.55), Inches(5.65), Inches(5.0), Inches(1.2),
            [(0, "活跃度较高，新增发电侧主体参与"), (0, "是中长期套利的重要市场，更注重“价”")], size=12.5)
footer(s, 8, TOTAL)

# ---------- Slide 9: 多日直接交易 — 策略 ----------
s = add_slide()
set_bg(s)
title_bar(s, "多日直接交易 — 交易策略", "07 Strategy")
add_rect(s, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.85), LIGHT)
add_text(s, Inches(0.85), Inches(1.62), Inches(11.6), Inches(0.6),
          "盈亏线 = 现货价格_预测；现货价格_预测 > 当前报价 → 买入，反之 → 卖出",
          size=14.5, color=NAVY, bold=True)

items = [
    (0, "集中竞价阶段"),
    (1, "结合已有持仓（含上午合同转让结果）规划当前市场可交易总量"),
    (1, "例：持仓50%+月竞10%，预知实时结算电量20%更便宜，则可交易空间 = 100%−50%−10%−20% = 20%"),
    (1, "粗判集中/滚动价格趋势，按个人偏好分配集中竞价电量（如拿出80%可交易量），竞价电价略低于预测价"),
    (0, "滚动撮合阶段"),
    (1, "成交量小，部分交易员只盯关键点位或直接放弃该阶段"),
    (1, "基于已有出清量，挂出比预测价略低的报价"),
    (1, "其余经验与合同电量转让市场一致"),
]
add_bullets(s, Inches(0.6), Inches(2.55), Inches(12.1), Inches(4.3), items, size=14.5)
footer(s, 9, TOTAL)

# ---------- Slide 10: 月度策略 + 补充 ----------
s = add_slide()
set_bg(s)
title_bar(s, "月度策略 / 全局套利构想 / 补充内容", "08-09 Outlook")

add_text(s, Inches(0.6), Inches(1.35), Inches(6), Inches(0.35), "月度策略（现状）", size=15, bold=True, color=NAVY)
add_bullets(s, Inches(0.6), Inches(1.75), Inches(5.9), Inches(1.3),
            [(0, "无法准确预测本月发电侧实时市场加权均价"),
             (0, "月初观望、月中启动并逐渐放量；放量比例人工决策"),
             (0, "原则：保证月末考核电量基本为0，但不绝对")], size=13)

add_text(s, Inches(0.6), Inches(3.15), Inches(8), Inches(0.35), "全局套利构想【待讨论】", size=15, bold=True, color=ORANGE)
add_bullets(s, Inches(0.6), Inches(3.55), Inches(5.9), Inches(2.6),
            [(0, "三层架构：月度分布 → 多日分布 → 单点策略"),
             (1, "月初/月内滚动判断全月电价趋势，确定多日电能量分布"),
             (1, "单次交易判断2~5天趋势，确定各时段交易电量分布"),
             (1, "单时刻点按盈亏概率报量报价"),
             (0, "需实时动态调整：月初与月末现货价差可达200–300元/MWh"),
             (0, "策略触发须明确告知交易员，由其决策是否执行")], size=12.5)

add_rect(s, Inches(6.9), Inches(1.35), Inches(5.85), Inches(4.85), LIGHT)
add_text(s, Inches(7.15), Inches(1.55), Inches(5.3), Inches(0.35), "补充内容 / 待完善方向", size=15, bold=True, color=NAVY)
add_bullets(s, Inches(7.15), Inches(2.0), Inches(5.3), Inches(4.0),
            [(0, "盯盘工具需求"),
             (1, "滚动撮合阶段对手方可能误报极低/高价，需工具辅助“抢单”"),
             (0, "交易水平评价体系"),
             (1, "交易品种多、频率高，目前缺乏明确的交易水平评价指标，后续需补充")], size=13)
footer(s, 10, TOTAL)

prs.save("/home/user/ppt/中长期交易基本说明与业务逻辑.pptx")
print("saved")
